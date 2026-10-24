"""Pruebas del paso 4: madurez ("lista mañana", dudosas) y reporte Excel + JSON."""

import json

import pytest
from openpyxl import load_workbook

import configuracion
from navegacion.plano import armar_plano
from percepcion import madurez
from percepcion.dedup import Rosa
from reporte import archivos
from reporte.tabla import COLUMNAS_CUADRANTES, COLUMNAS_LISTAS, armar_reporte, prioridad


@pytest.fixture(scope="module")
def cfg():
    return configuracion.cargar()


def probs(inm, opt):
    return {"inmaduro": inm, "optimo": opt, "pasado": round(1 - inm - opt, 3)}


def test_ejemplo_sin_historia_del_plan(cfg):
    # p(óptimo) = 0,45 y tamaño_rel = 0,92 → puntaje 0,64 → lista mañana (sección 2.4)
    m = madurez.decidir("inmaduro", True, probs(0.53, 0.45), 0.92 * 5.0, cfg)
    assert m.estado == "lista_manana" and m.puntaje == pytest.approx(0.638)


def test_ejemplo_con_historia_del_plan(cfg):
    # ayer 4,0 cm, hoy 4,5 cm, punto 3 = 5,0 cm → falta 1 día → lista mañana
    assert madurez.dias_faltan(4.5, 4.0, 5.0) == pytest.approx(1.0)
    m = madurez.decidir("inmaduro", True, probs(0.6, 0.35), 4.5, cfg, diametro_ayer_cm=4.0)
    assert m.estado == "lista_manana" and m.dias_faltan == pytest.approx(1.0)
    # si no creció, no se puede predecir con historia y se usa el puntaje
    assert madurez.dias_faltan(4.5, 4.6, 5.0) is None


def test_inmadura_pequena_o_sin_p3_no_es_lista_manana(cfg):
    assert madurez.decidir("inmaduro", True, probs(0.9, 0.08), 4.9, cfg).estado == "inmaduro"
    # tamaño_rel 0,7 < 0,85: el tamaño no suma y el puntaje no alcanza
    assert madurez.decidir("inmaduro", True, probs(0.55, 0.44), 3.5, cfg).estado == "inmaduro"


def test_dudosas(cfg):
    assert madurez.decidir("optimo", False, probs(0.3, 0.55), 5.0, cfg).estado == "dudosa"
    assert madurez.decidir("pasado", False, {"inmaduro": 0.2, "optimo": 0.3, "pasado": 0.5}, 7.0, cfg).estado == "dudosa"
    # duda solo entre inmaduro y óptimo: se decide con las reglas de "lista mañana"
    assert madurez.decidir("inmaduro", False, probs(0.52, 0.45), 4.6, cfg).estado == "lista_manana"
    assert madurez.decidir("optimo", True, probs(0.1, 0.85), 5.0, cfg).estado == "optimo"


def test_prioridad():
    assert prioridad(3, 0, 3) == "alta"
    assert prioridad(1, 0, 3) == prioridad(0, 2, 3) == "media"
    assert prioridad(0, 0, 3) == "baja"


def rosa(cama, tramo, metro, clase, conf=0.9, opt=None, diam=5.0, rid=None):
    p = {"inmaduro": 0.05, "optimo": 0.05, "pasado": 0.05}
    p[clase] = conf
    if opt is not None:
        p["optimo"] = opt
    cid = f"INV1-C{cama:02d}-T{tramo}"
    return Rosa(rid or f"{cid}-{metro}", cid, clase, conf, conf >= 0.6, p, cama, tramo, "izq", metro, 1.5,
                metro, 2.9, diam, 1.2, "fotos/a.png")


@pytest.fixture(scope="module")
def reporte(cfg):
    rosas = [rosa(2, 1, 1.0, "optimo"), rosa(2, 1, 2.0, "optimo"), rosa(2, 1, 3.0, "optimo"),
             rosa(2, 2, 6.0, "inmaduro", conf=0.53, opt=0.45, diam=4.6),
             rosa(1, 2, 7.0, "pasado"), rosa(1, 1, 2.0, "optimo", conf=0.5)]
    return armar_reporte(rosas, armar_plano(cfg), {1, 2}, cfg, {
        "id_ronda": "INV1_2026-10-24_220000", "inicio": "2026-10-24 22:00:00", "fin": "2026-10-24 23:41:00",
        "fechas_fotos": {"fotos/a.png": "2026-10-24 22:05:00"}, "ronda": [("fotos tomadas", 36)]})


def test_tabla_de_cuadrantes(reporte):
    cuad = {c["cuadrante_id"]: c for c in reporte.cuadrantes}
    assert len(cuad) == 4                                    # 2 camas × 2 tramos de 5 m, aunque estén vacíos
    c = cuad["INV1-C02-T1"]
    assert (c["optimo"], c["total"], c["prioridad"], c["orden"]) == (3, 3, "alta", 1)
    assert (cuad["INV1-C02-T2"]["lista_manana"], cuad["INV1-C02-T2"]["prioridad"]) == (1, "media")
    assert cuad["INV1-C01-T1"]["dudosa"] == 1 and cuad["INV1-C01-T2"]["prioridad"] == "baja"
    assert sorted(c["orden"] for c in reporte.cuadrantes) == [1, 2, 3, 4]
    assert (c["desde_m"], c["hasta_m"]) == (0, 5)


def test_rosas_listas_y_dudosas(reporte):
    estados = [f["estado"] for f in reporte.rosas_listas]
    assert estados == ["optimo"] * 3 + ["lista_manana"]      # las de hoy primero
    assert list(reporte.rosas_listas[0]) == COLUMNAS_LISTAS
    assert reporte.rosas_listas[0]["fecha_hora"] == "2026-10-24 22:05:00"
    assert len(reporte.dudosas) == 1 and len(reporte.todas) == 6


def test_excel_y_json_con_los_mismos_datos(reporte, tmp_path):
    xlsx, js = archivos.guardar(reporte, tmp_path, avisar=lambda s: None)
    assert xlsx.name == "reporte_INV1_2026-10-24.xlsx" and js.name == "reporte_INV1_2026-10-24.json"
    wb = load_workbook(xlsx)
    assert wb.sheetnames == ["Resumen_cuadrantes", "Rosas_listas", "Todas_las_rosas", "Dudosas", "Ronda", "Diccionario"]
    hoja = wb["Resumen_cuadrantes"]
    assert [c.value for c in hoja[1]] == COLUMNAS_CUADRANTES
    filas_excel = {r[0]: r for r in hoja.iter_rows(min_row=2, values_only=True)}
    datos = json.loads(js.read_text(encoding="utf-8"))
    for c in datos["cuadrantes"]:
        assert list(filas_excel[c["cuadrante_id"]]) == [c[k] for k in COLUMNAS_CUADRANTES]
    assert datos["id_ronda"] == "INV1_2026-10-24_220000" and datos["ronda"]["fotos tomadas"] == 36


def test_nunca_sobreescribe(reporte, tmp_path):
    primero = archivos.guardar(reporte, tmp_path, avisar=lambda s: None)
    segundo = archivos.guardar(reporte, tmp_path, avisar=lambda s: None)
    assert [p.name for p in segundo] == ["reporte_INV1_2026-10-24_v2.xlsx", "reporte_INV1_2026-10-24_v2.json"]
    assert all(p.exists() for p in primero + segundo)


def test_no_repite_el_nombre_de_uno_ya_enviado(reporte, tmp_path):
    salida, enviados = tmp_path / "salida", tmp_path / "enviados"
    enviados.mkdir()
    (enviados / "reporte_INV1_2026-10-24.xlsx").write_bytes(b"ya llego al dashboard")
    rutas = archivos.guardar(reporte, salida, avisar=lambda s: None, ya_enviados=enviados)
    assert [p.name for p in rutas] == ["reporte_INV1_2026-10-24_v2.xlsx", "reporte_INV1_2026-10-24_v2.json"]


def test_si_esta_abierto_usa_otro_nombre(reporte, tmp_path, monkeypatch):
    from openpyxl import Workbook
    original = Workbook.save
    def bloqueado(self, ruta):
        if not str(ruta).endswith("_v2.xlsx"):
            raise PermissionError("abierto en Excel")
        return original(self, ruta)
    monkeypatch.setattr(Workbook, "save", bloqueado)
    rutas = archivos.guardar(reporte, tmp_path, avisar=lambda s: None)
    assert rutas[0].name.endswith("_v2.xlsx") and rutas[0].exists()
