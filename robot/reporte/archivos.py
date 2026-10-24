"""Reporte de la ronda en Excel y en JSON (secciones 10.1 a 10.4 del plan).

- Un archivo nuevo por noche: reporte_INV1_2026-10-24.xlsx y .json en datos/salida/.
- Nunca se sobreescribe ni se cae: si ya existe (o está abierto en Excel), se
  guarda como _v2, _v3…
- Los colores son solo visuales (los mismos del juego de la tablet); el
  dashboard lee los datos.
"""

import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from reporte.tabla import (COLUMNAS_CUADRANTES, COLUMNAS_DUDOSAS, COLUMNAS_LISTAS, COLUMNAS_TODAS, Reporte)

# colores de app-tablet/src/data/categorias.ts, aclarados para que se lea el texto
COLORES = {
    "optimo": "F2C4BF",        # #C0392B
    "lista_manana": "F9E0C7",
    "inmaduro": "D6E4C8",      # #5B7F3A
    "pasado": "E3D2C2",        # #8E5A2B
    "dudosa": "E0E0E0",
    "alta": "F2C4BF",
    "media": "F9E0C7",
    "baja": "FFFFFF",
}
MAX_VERSIONES = 50


def _hoja(wb: Workbook, titulo: str, columnas: list[str], filas: list[dict], colorear: str | None = None):
    ws = wb.create_sheet(titulo)
    ws.append(columnas)
    for celda in ws[1]:
        celda.font = Font(bold=True)
    for fila in filas:
        ws.append([fila.get(c) for c in columnas])
        if colorear and fila.get(colorear) in COLORES:
            relleno = PatternFill("solid", fgColor=COLORES[fila[colorear]])
            for celda in ws[ws.max_row]:
                celda.fill = relleno
    ws.auto_filter.ref = ws.dimensions
    ws.freeze_panes = "A2"
    for i, c in enumerate(columnas, start=1):
        largo = max([len(str(c))] + [len(str(f.get(c, ""))) for f in filas])
        ws.column_dimensions[get_column_letter(i)].width = min(largo + 2, 60)
    return ws


def libro(rep: Reporte) -> Workbook:
    wb = Workbook()
    wb.remove(wb.active)
    _hoja(wb, "Resumen_cuadrantes", COLUMNAS_CUADRANTES, sorted(rep.cuadrantes, key=lambda f: f["orden"]), "prioridad")
    _hoja(wb, "Rosas_listas", COLUMNAS_LISTAS, rep.rosas_listas, "estado")
    _hoja(wb, "Todas_las_rosas", COLUMNAS_TODAS, rep.todas, "estado")
    _hoja(wb, "Dudosas", COLUMNAS_DUDOSAS, rep.dudosas)
    _hoja(wb, "Ronda", ["campo", "valor"], [{"campo": k, "valor": str(v)} for k, v in rep.ronda])
    _hoja(wb, "Diccionario", ["columna", "significado"], [{"columna": k, "significado": v} for k, v in rep.diccionario])
    return wb


def a_json(rep: Reporte) -> dict:
    return {
        "id_ronda": rep.id_ronda,
        "invernadero": rep.invernadero,
        "fecha": rep.fecha,
        "inicio": rep.inicio,
        "fin": rep.fin,
        "cuadrantes": rep.cuadrantes,
        "rosas_listas": rep.rosas_listas,
        "dudosas": rep.dudosas,
        "obstaculos": rep.obstaculos,
        "ronda": {k: v for k, v in rep.ronda},
    }


def _nombre(base: str, version: int) -> str:
    return base if version == 1 else f"{base}_v{version}"


def guardar(rep: Reporte, carpeta: Path, formatos=("excel", "json"), avisar=print,
            ya_enviados: Path | None = None) -> list[Path]:
    """Guarda el reporte con un nombre libre. Devuelve las rutas escritas.

    `ya_enviados`: carpeta de los reportes que ya salieron (datos/enviados/). Sus
    nombres tampoco se repiten, para no pisar en el dashboard uno que ya llegó.
    """
    carpeta.mkdir(parents=True, exist_ok=True)
    base = f"reporte_{rep.invernadero}_{rep.fecha}"
    extensiones = {"excel": ".xlsx", "json": ".json"}
    for version in range(1, MAX_VERSIONES + 1):
        nombre = _nombre(base, version)
        rutas = {f: carpeta / f"{nombre}{extensiones[f]}" for f in formatos}
        usados = list(rutas.values())
        if ya_enviados is not None:
            usados += [ya_enviados / r.name for r in rutas.values()]
        if any(r.exists() for r in usados):
            continue
        try:
            if "excel" in rutas:
                libro(rep).save(rutas["excel"])
            if "json" in rutas:
                with open(rutas["json"], "w", encoding="utf-8") as f:
                    json.dump(a_json(rep), f, ensure_ascii=False, indent=1)
        except PermissionError:       # abierto en Excel o sin permiso: se prueba el siguiente nombre
            for r in rutas.values():
                try:
                    r.unlink(missing_ok=True)
                except PermissionError:
                    pass
            continue
        if version > 1:
            avisar(f"  aviso: ya había un reporte de esta noche; se guardó como {nombre}")
        return list(rutas.values())
    raise RuntimeError(f"no se pudo guardar el reporte en {carpeta} (más de {MAX_VERSIONES} versiones)")
